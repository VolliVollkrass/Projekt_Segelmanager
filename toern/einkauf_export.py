"""Einkaufsliste als PDF (zum Ausdrucken) und als Excel-Datei.

Beide Exporte zeigen die aktive Liste eines Bootes — archivierte Posten
(Kaufhistorie) bleiben außen vor. Gruppiert wird nach Kategorie, in derselben
Reihenfolge wie im Boot-Dashboard. Die Einkäufer-Spalte kommt aus „Aufteilen"
und bleibt leer, solange die Liste niemandem zugeteilt ist.
"""
import os
from io import BytesIO

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

from boote.models import Boot
from logistik.models import EinkaufslistenEintrag
from .models import Teilnahme, Toern

LOGO_PATH = os.path.join(settings.BASE_DIR, 'static', 'medien', 'Logo_Meer_erleben.png')

PRIMARY = colors.HexColor('#1e3a5f')
SECONDARY = colors.HexColor('#0D9488')
GRAY = colors.HexColor('#9ca3af')
GRAY_LIGHT = colors.HexColor('#d1d5db')


def _daten(request, toern_id, boot_id):
    """Gemeinsame Vorarbeit beider Exporte: Rechte prüfen, Posten gruppieren."""
    from .views import EINKAUF_KATEGORIE_LABEL, EINKAUF_KATEGORIE_ORDER

    toern = get_object_or_404(Toern, id=toern_id)
    boot = get_object_or_404(Boot, id=boot_id)

    # Gleiche Rechte wie beim Bearbeiten der Liste (siehe einkaufsliste_add)
    if not Teilnahme.objects.filter(user=request.user, toern=toern).exists() \
            and request.user != toern.anbieter:
        raise PermissionDenied

    eintraege = (
        EinkaufslistenEintrag.objects
        .filter(boot=boot, toern=toern, archiviert=False)
        .select_related('einkaufer__user')
        .order_by('name')
    )

    by_kat = {}
    for e in eintraege:
        by_kat.setdefault(e.kategorie, []).append(e)

    gruppen = [
        (EINKAUF_KATEGORIE_LABEL.get(kat, 'Sonstiges'), by_kat[kat])
        for kat in EINKAUF_KATEGORIE_ORDER
        if by_kat.get(kat)
    ]
    # Kategorien, die (noch) nicht in der Reihenfolge stehen, hinten anhängen
    for kat, items in by_kat.items():
        if kat not in EINKAUF_KATEGORIE_ORDER:
            gruppen.append((EINKAUF_KATEGORIE_LABEL.get(kat, 'Sonstiges'), items))

    return toern, boot, gruppen


def _einkaufer_name(eintrag):
    if not eintrag.einkaufer:
        return ''
    u = eintrag.einkaufer.user
    return f"{u.first_name} {u.last_name}".strip()


def _dateiname(boot, endung):
    safe = ''.join(c if c.isalnum() or c in '-_' else '_' for c in boot.name)
    return f"einkaufsliste_{safe}.{endung}"


# ══════════════════════════════════════════════════════════════════════════
# PDF
# ══════════════════════════════════════════════════════════════════════════
@login_required
def einkaufsliste_pdf(request, toern_id, boot_id):
    toern, boot, gruppen = _daten(request, toern_id, boot_id)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=18 * mm,
        title=f"Einkaufsliste {boot.name}",
    )

    titel_s = ParagraphStyle('Titel', fontSize=15, leading=19,
                             fontName='Helvetica-Bold', textColor=PRIMARY)
    sub_s = ParagraphStyle('Sub', fontSize=9, leading=12, textColor=GRAY)
    kat_s = ParagraphStyle('Kat', fontSize=10, leading=13,
                           fontName='Helvetica-Bold', textColor=SECONDARY)
    zelle_s = ParagraphStyle('Zelle', fontSize=9.5, leading=12.5)
    zelle_durch_s = ParagraphStyle('ZelleDurch', parent=zelle_s, textColor=GRAY)
    kopf_s = ParagraphStyle('Kopf', fontSize=8, leading=11,
                            fontName='Helvetica-Bold', textColor=colors.white)

    elemente = []

    # ─── Kopf: Titel links, Logo rechts ───
    kopf_links = [
        Paragraph(f"Einkaufsliste — {boot.name}", titel_s),
        Spacer(1, 2 * mm),
        Paragraph(
            f"{toern.titel} &nbsp;|&nbsp; "
            f"{toern.startdatum.strftime('%d.%m.%Y')} – {toern.enddatum.strftime('%d.%m.%Y')}",
            sub_s,
        ),
    ]
    logo = Image(LOGO_PATH, width=30 * mm, height=30 * mm, kind='proportional') \
        if os.path.exists(LOGO_PATH) else ''
    kopf = Table([[kopf_links, logo]], colWidths=[135 * mm, 45 * mm])
    kopf.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (1, 0), (1, 0), 'RIGHT'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
    ]))
    elemente.append(kopf)
    elemente.append(Spacer(1, 6 * mm))

    if not gruppen:
        elemente.append(Paragraph('Die Einkaufsliste ist leer.', zelle_s))
    else:
        SPALTEN = [10 * mm, 86 * mm, 38 * mm, 46 * mm]

        for label, items in gruppen:
            daten = [[
                '',
                Paragraph('Artikel', kopf_s),
                Paragraph('Menge', kopf_s),
                Paragraph('Einkäufer', kopf_s),
            ]]
            for e in items:
                stil = zelle_durch_s if e.erledigt else zelle_s
                name = f"<strike>{e.name}</strike>" if e.erledigt else e.name
                daten.append([
                    '✗' if e.erledigt else '',
                    Paragraph(name, stil),
                    Paragraph(e.menge or '', stil),
                    Paragraph(_einkaufer_name(e), stil),
                ])

            tabelle = Table(daten, colWidths=SPALTEN, repeatRows=1)
            tabelle.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), PRIMARY),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('ALIGN', (0, 0), (0, -1), 'CENTER'),
                ('FONTSIZE', (0, 1), (0, -1), 10),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                # Ankreuzkästchen in der ersten Spalte
                ('GRID', (0, 1), (0, -1), 0.6, GRAY),
                ('LINEBELOW', (1, 1), (-1, -2), 0.4, GRAY_LIGHT),
            ]))

            elemente.append(KeepTogether([
                Paragraph(f"{label} <font color='#9ca3af'>({len(items)})</font>", kat_s),
                Spacer(1, 1.5 * mm),
                tabelle,
                Spacer(1, 5 * mm),
            ]))

    def fusszeile(canvas, _doc):
        canvas.saveState()
        breite, _hoehe = A4
        canvas.setStrokeColor(GRAY_LIGHT)
        canvas.setLineWidth(0.5)
        canvas.line(15 * mm, 13 * mm, breite - 15 * mm, 13 * mm)
        canvas.setFont('Helvetica', 7)
        canvas.setFillColor(GRAY)
        canvas.drawString(
            15 * mm, 9 * mm,
            f"Einkaufsliste {boot.name} · Stand "
            f"{timezone.localtime().strftime('%d.%m.%Y %H:%M')} Uhr",
        )
        canvas.drawRightString(breite - 15 * mm, 9 * mm, f"Seite {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(elemente, onFirstPage=fusszeile, onLaterPages=fusszeile)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{_dateiname(boot, "pdf")}"'
    return response


# ══════════════════════════════════════════════════════════════════════════
# EXCEL
# ══════════════════════════════════════════════════════════════════════════
@login_required
def einkaufsliste_xlsx(request, toern_id, boot_id):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    toern, boot, gruppen = _daten(request, toern_id, boot_id)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Einkaufsliste'
    ws.sheet_view.showGridLines = False

    PRIMARY_HEX = '1E3A5F'
    TEAL_HEX = '0D9488'
    LIGHT = 'E2E8F0'

    kopf_font = Font(bold=True, color='FFFFFF')
    kopf_fill = PatternFill('solid', fgColor=PRIMARY_HEX)
    kat_font = Font(bold=True, color='FFFFFF')
    kat_fill = PatternFill('solid', fgColor=TEAL_HEX)
    thin = Side(style='thin', color=LIGHT)
    rahmen = Border(bottom=thin)
    grau = Font(color='9CA3AF')

    # Spalte A bleibt Rand; Inhalt ab Spalte B — wie beim Abrechnungs-Export
    ws.append([])
    ws.append([None, f"Einkaufsliste – {boot.name}"])
    ws['B2'].font = Font(bold=True, size=14, color=PRIMARY_HEX)
    ws.append([None, f"{toern.titel} · {toern.startdatum.strftime('%d.%m.%Y')} – "
                     f"{toern.enddatum.strftime('%d.%m.%Y')}"])
    ws['B3'].font = Font(italic=True, color='64748B')
    ws.append([None, f"Stand {timezone.localtime().strftime('%d.%m.%Y %H:%M')} Uhr"])
    ws['B4'].font = Font(italic=True, color='64748B')
    ws.append([])

    spalten = ['Gekauft', 'Artikel', 'Menge', 'Einkäufer']
    ws.append([None] + spalten)
    # Zeilennummer NACH dem append lesen: ein leeres append([]) schiebt zwar den
    # Schreib-Cursor weiter, erhöht aber max_row nicht — vorher gerechnet läge
    # die Formatierung eine Zeile zu hoch.
    kopf_zeile = ws.max_row
    for spalte in range(2, 2 + len(spalten)):
        zelle = ws.cell(row=kopf_zeile, column=spalte)
        zelle.font = kopf_font
        zelle.fill = kopf_fill
        zelle.alignment = Alignment(horizontal='center')

    for label, items in gruppen:
        ws.append([None, label, '', '', ''])
        zeile = ws.max_row
        for spalte in range(2, 6):
            ws.cell(row=zeile, column=spalte).fill = kat_fill
            ws.cell(row=zeile, column=spalte).font = kat_font

        for e in items:
            ws.append([
                None,
                'x' if e.erledigt else '',
                e.name,
                e.menge or '',
                _einkaufer_name(e),
            ])
            zeile = ws.max_row
            ws.cell(row=zeile, column=2).alignment = Alignment(horizontal='center')
            for spalte in range(2, 6):
                ws.cell(row=zeile, column=spalte).border = rahmen
                if e.erledigt:
                    ws.cell(row=zeile, column=spalte).font = grau

    if not gruppen:
        ws.append([None, '', 'Die Einkaufsliste ist leer.', '', ''])

    for spalte, breite in zip(range(2, 6), [10, 40, 18, 24]):
        ws.column_dimensions[get_column_letter(spalte)].width = breite
    ws.column_dimensions['A'].width = 3
    ws.freeze_panes = ws.cell(row=kopf_zeile + 1, column=1)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{_dateiname(boot, "xlsx")}"'
    return response
